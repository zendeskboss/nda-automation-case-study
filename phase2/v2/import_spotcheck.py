"""Read the lawyer's filled spot-check sheet, record each answer in the challenge log, and set the default policy.

  python -X utf8 phase2/v2/import_spotcheck.py phase2/v2/spotcheck/SPOTCHECK_v2.1.xlsx

Rule agreed with the user: v2.1 becomes the default only when all rows are filled in (verdict, reviewer, date)
and no row is overturned ('deviation'). Otherwise v2.0 stays the default. The decision is written to
default_policy.json, which run.py reads.
"""
import datetime, hashlib, json, os, sys, time
from openpyxl import load_workbook

V2 = os.path.dirname(os.path.abspath(__file__))
LOG = os.path.join(V2, "challenges.jsonl")
DEFAULT = os.path.join(V2, "default_policy.json")
VERDICTS = {"matches", "acceptable", "deviation"}

def main(path, log=LOG, default=DEFAULT):
    ws = load_workbook(path, data_only=True)["Spot check"]
    hdr = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(hdr)}
    rows, problems = [], []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[col["Contract"]] is None: continue
        v = (r[col["Lawyer verdict"]] or "").strip().lower()
        by = (r[col["Reviewed by"]] or "").strip() if isinstance(r[col["Reviewed by"]], str) else r[col["Reviewed by"]]
        d = r[col["Date (YYYY-MM-DD)"]]
        if isinstance(d, (datetime.date, datetime.datetime)): d = d.strftime("%Y-%m-%d")
        item = dict(contract=r[col["Contract"]], clause=str(r[col["Clause"]]), rule=r[col["Rule"]], tool_route=r[col["v2.1 result"]],
                    lawyer_verdict=v, want_in_queue=r[col["Want clauses like this in my queue?"]], note=r[col["Note"]] or "", by=by, date=d)
        if v not in VERDICTS: problems.append("%s %s: verdict '%s' is not matches, acceptable or deviation" % (item["contract"], item["clause"], v))
        if not by or not d: problems.append("%s %s: reviewer or date is missing" % (item["contract"], item["clause"]))
        rows.append(item)
    if problems:
        print("The sheet is not complete. Nothing was recorded.\n- " + "\n- ".join(problems)); return 2
    sheet_hash = hashlib.sha256(open(path, "rb").read()).hexdigest()[:16]
    with open(log, "a", encoding="utf-8") as fh:
        for it in rows:
            fh.write(json.dumps(dict(it, time=time.strftime("%Y-%m-%dT%H:%M:%S"), source="spot-check v2.1", sheet_sha256=sheet_hash,
                                     run="r1", policy="v2.1", tool_verdict=it["tool_route"]), ensure_ascii=False, default=str) + "\n")
    over = [it for it in rows if it["lawyer_verdict"] == "deviation"]
    decision = dict(default="v2.0" if over else "v2.1", decided=time.strftime("%Y-%m-%dT%H:%M:%S"), source=os.path.basename(path),
                    sheet_sha256=sheet_hash, rows=len(rows), overturned=["%s %s %s" % (o["contract"], o["clause"], o["rule"]) for o in over],
                    reviewers=sorted({str(it["by"]) for it in rows}))
    json.dump(decision, open(default, "w", encoding="utf-8"), indent=1)
    print("Recorded %d answers in %s." % (len(rows), os.path.relpath(log)))
    print("Default policy: %s%s" % (decision["default"], "" if not over else " (overturned: %s)" % ", ".join(decision["overturned"])))
    return 0

if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
