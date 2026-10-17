"""Build the lawyer spot-check sheet for policy v2.1: the 8 clauses that v1 and v2.0 sent to a lawyer and v2.1 passes.

  python -X utf8 phase2/v2/make_spotcheck.py      writes phase2/v2/spotcheck/SPOTCHECK_v2.1.xlsx

The lawyer fills in the yellow cells. Then `import_spotcheck.py` records the answers and sets the default policy.
"""
import json, os, re, sys
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import engine, review, jev_test  # noqa: E402

V2, ROOT = engine.V2, engine.ROOT
ITEMS = [("C14", "6.1", "PB07"), ("C15", "10.1", "PB19"), ("C24", "1.1.7", "PB05"), ("C27", "1.1.4", "PB05"),
         ("C29", "4.3", "PB06"), ("C30", "1.1.6", "PB05"), ("C30", "2.1", "PB20"), ("C30", "9.4", "PB19")]
OUT = os.path.join(V2, "spotcheck"); os.makedirs(OUT, exist_ok=True)
PATH = os.path.join(OUT, "SPOTCHECK_v2.1.xlsx")

pb = review.load_playbook()
rules, _ = engine.catalogue(pb)
sev = {r["id"]: r["severity"] for r in pb["rules"]}

def mc(c):
    m = re.match(r"\s*(\d+(?:\.\d+)?)", str(c or "")); return m.group(1) if m else str(c)

rows = []
for cid, cl, rid in ITEMS:
    text = open(os.path.join(ROOT, "contracts", cid + ".md"), encoding="utf-8").read()
    f21 = next(f for f in json.load(open(os.path.join(V2, "runs", "r1", cid + ".raw.v2.1.json"), encoding="utf-8"))["findings"]
               if f["clause_no"] == cl and f["rule_id"] == rid)
    r21 = next(f for f in json.load(open(os.path.join(V2, "runs", "r1", cid + ".routed.v2.1.json"), encoding="utf-8"))["findings"]
               if f["clause_no"] == cl and f["rule_id"] == rid)
    v1 = [f for f in review.route(json.load(open(os.path.join(ROOT, "replay", cid + ".raw.json"), encoding="utf-8")), pb, text=text)["findings"]
          if f["rule_id"] == rid and mc(f["clause_no"]) == mc(cl) and f["verdict"] in ("escalate", "deviation_draft")]
    pr = f21["evidence"]["answers"]["position_check"]["probabilities"]
    assert jev_test.quote_found(f21["quote"], text), (cid, cl)
    rows.append(dict(cid=cid, cl=cl, rid=rid, topic=rules[rid]["topic"], sev=sev[rid],
                     preferred=rules[rid]["preferred_position"],
                     fallbacks="\n".join("• " + a for a in rules[rid]["acceptable_fallbacks"]) or "(none listed)",
                     clause=f21["quote"], v21=r21["verdict"],
                     probs="pass %.2f (matches %.2f, fallback %.2f) · breach %.2f · unclear %.2f" % (
                         pr["matches_preferred"] + pr["meets_fallback"], pr["matches_preferred"], pr["meets_fallback"], pr["breaches"], pr["unclear"]),
                     v1="; ".join("%s (clause %s): %s" % (f["ai_verdict"], f["clause_no"], f["why"]) for f in v1)))

ARIAL = "Arial"
H = Font(name=ARIAL, bold=True, color="FFFFFF"); B = Font(name=ARIAL); BB = Font(name=ARIAL, bold=True)
HFILL = PatternFill("solid", fgColor="1F4E5A"); IN = PatternFill("solid", fgColor="FFFF00"); GREY = PatternFill("solid", fgColor="EEF1F2")
thin = Side(style="thin", color="B7BEC2"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
# ---------- Instructions ----------
ins = wb.active; ins.title = "Instructions"
lines = [
    ("NDA review v2.1: lawyer spot-check", BB),
    ("", B),
    ("Why you are asked", BB),
    ("Policy v2.1 reduces the review load. It passes 8 clauses that the earlier versions sent to a lawyer. "
     "Please confirm or overturn each one. v2.1 becomes the default only after this check.", B),
    ("", B),
    ("What to do", BB),
    ("1. Open the sheet 'Spot check'. Each row is one clause.", B),
    ("2. Read the clause text (column I) against the playbook position (columns G and H).", B),
    ("3. Fill in the yellow cells only: your verdict (M), whether you want clauses like this in your queue (N), a short note (O), your name (P) and the date (Q).", B),
    ("4. Do not change any other cell. Columns J to L show what the tool and the earlier review said, for reference only.", B),
    ("5. Save the file and return it. The import script records your answers in the challenge log.", B),
    ("", B),
    ("Verdict meanings (column M)", BB),
    ("matches: the clause gives the playbook's preferred position.", B),
    ("acceptable: the clause differs, but a listed acceptable fallback fits it.", B),
    ("deviation: the clause gives a result that the playbook does not allow. This overturns the tool's pass.", B),
    ("", B),
    ("What happens next", BB),
    ("All 8 filled in and no 'deviation': v2.1 becomes the default policy.", B),
    ("Any 'deviation': v2.0 stays the default, and the overturned rows go into the threshold review.", B),
    ("Column N is for tuning only. It does not block the switch.", B),
    ("", B),
    ("Example of a filled row (format only, not a real clause)", BB),
]
for i, (t, f) in enumerate(lines, 1):
    c = ins.cell(row=i, column=1, value=t); c.font = f; c.alignment = Alignment(wrap_text=True, vertical="top")
ex_hdr = ["Lawyer verdict (M)", "Want these in my queue? (N)", "Note (O)", "Reviewed by (P)", "Date (Q)"]
ex_val = ["deviation", "Yes", "Adds contractors to Representatives without liability for them; outside PB05.", "J. Smith", "2026-10-20"]
r0 = len(lines) + 1
for j, (h, v) in enumerate(zip(ex_hdr, ex_val)):
    a = ins.cell(row=r0, column=1 + j, value=h); a.font = H; a.fill = HFILL; a.border = BOX; a.alignment = WRAP
    b = ins.cell(row=r0 + 1, column=1 + j, value=v); b.font = Font(name=ARIAL, italic=True); b.fill = GREY; b.border = BOX; b.alignment = WRAP
ins.column_dimensions["A"].width = 110
for col in "BCDE": ins.column_dimensions[col].width = 22

# ---------- Spot check ----------
ws = wb.create_sheet("Spot check")
hdr = ["#", "Contract", "Clause", "Rule", "Topic", "Severity", "Playbook: preferred position", "Playbook: acceptable fallbacks",
       "Clause text (as in the contract, with any tracked edits)", "v2.1 result", "Jev probabilities", "Earlier review (v1 LLM) said",
       "Lawyer verdict", "Want clauses like this in my queue?", "Note", "Reviewed by", "Date (YYYY-MM-DD)"]
widths = [4, 9, 8, 7, 22, 9, 34, 38, 70, 11, 30, 50, 15, 14, 40, 16, 13]
for j, (h, w) in enumerate(zip(hdr, widths), 1):
    c = ws.cell(row=1, column=j, value=h); c.font = H; c.fill = HFILL; c.border = BOX; c.alignment = WRAP
    ws.column_dimensions[c.column_letter].width = w
for i, r in enumerate(rows, 2):
    vals = [i - 1, r["cid"], r["cl"], r["rid"], r["topic"], r["sev"], r["preferred"], r["fallbacks"], r["clause"],
            r["v21"], r["probs"], r["v1"] or "(no matching v1 finding)"]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=i, column=j, value=v); c.font = B; c.border = BOX; c.alignment = WRAP
    for j in range(13, 18):
        c = ws.cell(row=i, column=j); c.fill = IN; c.border = BOX; c.font = B; c.alignment = WRAP
    ws.row_dimensions[i].height = 230
last = len(rows) + 1
dv1 = DataValidation(type="list", formula1='"matches,acceptable,deviation"', allow_blank=True,
                     error="Choose matches, acceptable or deviation.", errorTitle="Verdict")
dv2 = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
dv3 = DataValidation(type="date", operator="greaterThan", formula1="DATE(2026,1,1)", allow_blank=True, error="Enter a date.")
for dv, col in ((dv1, "M"), (dv2, "N"), (dv3, "Q")):
    ws.add_data_validation(dv); dv.add("%s2:%s%d" % (col, col, last))
for i in range(2, last + 1): ws.cell(row=i, column=17).number_format = "yyyy-mm-dd"
ws.freeze_panes = "B2"
ws["M1"].comment = Comment("Fill in: matches, acceptable or deviation. 'deviation' overturns the tool's pass.", "NDA review v2")
ws["J1"].comment = Comment("What policy v2.1 decided. 'matches' and 'acceptable' do not reach a lawyer.", "NDA review v2")

# ---------- Status (formulas) ----------
st = wb.create_sheet("Status")
spec = [
    ("Rows to check", "=COUNTA('Spot check'!B2:B%d)" % last),
    ("Verdicts filled in", "=COUNTA('Spot check'!M2:M%d)" % last),
    ("Rows with reviewer and date", "=SUMPRODUCT(('Spot check'!P2:P%d<>\"\")*('Spot check'!Q2:Q%d<>\"\"))" % (last, last)),
    ("Overturned (verdict = deviation)", "=COUNTIF('Spot check'!M2:M%d,\"deviation\")" % last),
    ("Want in queue = Yes", "=COUNTIF('Spot check'!N2:N%d,\"Yes\")" % last),
    ("Outcome", "=IF(OR(B3<B2,B4<B2),\"Waiting for the check\",IF(B5=0,\"Adopt v2.1 as the default\",\"Keep v2.0; review the overturned rows\"))"),
]
for i, (k, f) in enumerate(spec, 2):
    a = st.cell(row=i, column=1, value=k); a.font = BB if k == "Outcome" else B
    b = st.cell(row=i, column=2, value=f); b.font = BB if k == "Outcome" else B
st.cell(row=1, column=1, value="Check status").font = H; st.cell(row=1, column=1).fill = HFILL
st.cell(row=1, column=2, value="Value").font = H; st.cell(row=1, column=2).fill = HFILL
st.cell(row=9, column=1, value="Rule (from the agreed plan): v2.1 becomes the default when all rows are filled in and none is overturned.").font = Font(name=ARIAL, italic=True)
st.column_dimensions["A"].width = 40; st.column_dimensions["B"].width = 42

wb.active = 1
wb.calculation.fullCalcOnLoad = True  # Excel computes the Status formulas when the file opens
wb.save(PATH)
print(PATH, len(rows), "rows")
