"""Convert the lawyers' playbook worksheet into a draft playbook.

Run: python3 worksheet/worksheet_to_yaml.py worksheet/playbook_worksheet.xlsx worksheet/playbook.draft.yaml

The output has the same shape as playbook.yaml, so review.py can load it. It is a draft:
the playbook owner reads it, then freezes it as playbook.yaml with a new version number.
Rows with no preferred position are skipped and listed, so gaps are visible, not silent.
"""
import sys
import yaml
from openpyxl import load_workbook

LANES = {"AI drafts": "purple_ai", "Shadow": "purple_human", "Lawyer only": "orange"}
CATCH_ALL = dict(id="PB22", topic="Unrecognised clause (catch-all)", template_ref=[], risk_ids=[], severity="High",
                 lane="orange", preferred="n/a", acceptable=[],
                 deviation=["Any clause or obligation the tool cannot map to PB01-PB21 with confidence."],
                 draft_guidance=None)

def cell(v):
    return str(v).strip() if v is not None and str(v).strip() else None

def lines(v):
    return [x.strip() for x in str(v).splitlines() if x.strip()] if cell(v) else []

def ids(v):
    return [x.strip() for x in str(v).split(",") if x.strip()] if cell(v) else []

def suggest_lane(q1, q2, q3):
    # Same logic as the 'Suggested lane' formula in the sheet.
    if not (q1 and q2 and q3): return None
    if q3 == "No": return "Lawyer only"
    if q1 == "Yes": return "Shadow" if q2 == "Yes" else "Lawyer only"
    return "AI drafts" if q2 == "Yes" else "Lawyer only"

def convert(path):
    wb = load_workbook(path)
    warnings = []

    meta = {}
    for row in wb["1 Scope"].iter_rows(min_row=2, values_only=True):
        if cell(row[0]): meta[row[0]] = cell(row[2])
    meta["status"] = "draft"
    meta["severity_scale"] = {}
    for row in wb["2 Severity"].iter_rows(min_row=2, max_row=4, values_only=True):
        meta["severity_scale"][row[0]] = dict(cost_of_mistake=row[2], meaning=cell(row[1]))

    rules = []
    for row in wb["3 Clause positions"].iter_rows(min_row=2, values_only=True):
        rid, topic = cell(row[0]), cell(row[1])
        if not rid or not rid.startswith("PB"): continue
        preferred = cell(row[4])
        if not preferred:
            if topic and not topic.startswith("(new topic"): warnings.append("%s %s: no preferred position, skipped" % (rid, topic))
            continue
        lane_word = cell(row[12]) or suggest_lane(cell(row[8]), cell(row[9]), cell(row[10]))
        if lane_word not in LANES:
            warnings.append("%s: lane questions not answered, set to Lawyer only" % rid); lane_word = "Lawyer only"
        assumed = cell(row[15]) != "Yes"
        acceptable = [a + " (ASSUMED)" if assumed else a for a in lines(row[5])]
        if assumed: warnings.append("%s: not confirmed by a lawyer, fallbacks marked (ASSUMED)" % rid)
        rules.append(dict(id=rid, topic=topic, template_ref=ids(row[2]), risk_ids=ids(row[3]),
                          severity=cell(row[7]) or "High", lane=LANES[lane_word], preferred=preferred,
                          acceptable=acceptable, deviation=lines(row[6]), draft_guidance=cell(row[13]),
                          note_to_counterparty=cell(row[14])))
    rules.append(CATCH_ALL)

    requests = {}
    for row in wb["4 Known requests"].iter_rows(min_row=2, values_only=True):
        rq, topic, rule, answer, esc = (cell(x) for x in row[:5])
        if not rq or not topic: continue
        if not answer:
            warnings.append("%s %s: no standard answer, skipped" % (rq, topic)); continue
        requests[rq] = dict(topic=topic, rule=None if rule in (None, "None") else rule,
                            answer=answer + (" Escalate if: " + esc if esc else ""))
    requests["RQ12"] = dict(topic="Other / novel", rule="PB22", answer="Escalate.")

    return dict(meta=meta, rules=rules, known_requests=requests), warnings

if __name__ == "__main__":
    src, dst = sys.argv[1], sys.argv[2]
    pb, warnings = convert(src)
    with open(dst, "w", encoding="utf-8") as f:
        yaml.safe_dump(pb, f, sort_keys=False, allow_unicode=True, width=120)
    print("wrote %s: %d rules, %d known requests" % (dst, len(pb["rules"]), len(pb["known_requests"])))
    for w in warnings: print("  -", w)
